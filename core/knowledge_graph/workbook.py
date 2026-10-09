"""The workbook in which a case study is recorded, and its conversion.

A case study is entered in one Excel workbook, one sheet per kind of
record, so that it can be filled in without touching code or RDF:

    Project        what the case study is, its context and its BIM model
    Organisations  Organisation of each person
    People         who took part
    Activities     what was done, when, by whom, and in which role
    Documents      drawings, point clouds, publications, the BIM model...
    Units          structural units, and their floors
    Facades        each facade (or portion of a facade) and its masonry type
    Photos         the photographs of each facade
    Connections    documented connections between facades
    Openings       historic openings
    Vulnerabilities  the vulnerabilities of each facade wall (FMO), with evidence
    MasonryTypes   the seven MQI parameters of each masonry type, with the
                   evidence for each, its flags and its density
    Measured       measured properties to compare the derived ones with

read_workbook() turns it into the two records the graph is built from: the
case record and the masonry classification (the format of
resources/survey_data/*/masonry_classification.json). write_workbook()
does the reverse, and with no records writes the empty template.

Conventions in the cells: identifiers are short names without spaces
(they become the local names of the graph's IRIs); several values in one
cell are separated by ";"; roles are written "PersonId: Role"; variants
"parameter=value"; flags one per line. Dates are ISO (2023-06-01, or
2023-07-04T16:58:25Z for a time).
"""
import datetime

from core.ifc_processing import hmo_mqi

PARAMETERS = ["unit_dimensions", "unit_shape", "unit_material", "mortar_quality",
              "horizontal_joints", "vertical_joints", "wall_connections"]
CHOICES = {
    "unit_dimensions": list(hmo_mqi.UNIT_DIMENSIONS),
    "unit_shape": list(hmo_mqi.UNIT_SHAPE),
    "unit_material": list(hmo_mqi.UNIT_MATERIAL),
    "mortar_quality": list(hmo_mqi.MORTAR_QUALITY),
    "horizontal_joints": list(hmo_mqi.HORIZONTAL_JOINTS),
    "vertical_joints": list(hmo_mqi.VERTICAL_JOINTS),
    "wall_connections": list(hmo_mqi.WALL_CONNECTIONS),
    "document kind": ["Photo", "PointCloud", "Drawing", "Document", "Publication", "BuildingInformationModel"],
    "floor kind": ["TimberFloor", "Vault", "HorizontalStructure"],
    "connection kind": ["Connection", "Quoin", "Tie"],
    "yes/no": ["yes", "no"],
    "vulnerability": ["AbsenceOfTopChains", "AbsenceOfTopCurbstone", "AbsenceOfIntermediateChains",
                      "AbsenceOfIntermediateCurbstone", "DeformableFloors", "ExcessiveSlenderness",
                      "OpeningsNearIntersections", "PresenceOfHorizontalThrust"],
}

# sheet -> [(column, description, choices key or None)]
SHEETS = {
    "Project": [("key", "Field name (fixed)", None), ("value", "Value", None)],
    "Organisations": [("id", "Identifier", None), ("name", "Full name", None)],
    "People": [("id", "Identifier, e.g. MariaLauraLeonardi", None), ("name", "Full name", None),
               ("organisation", "Identifier of the organisation", None)],
    "Activities": [("id", "Identifier", None), ("label", "What was done", None),
                   ("date", "ISO date", None), ("start", "ISO date and time (optional)", None),
                   ("end", "ISO date and time (optional)", None),
                   ("people", "Identifiers of the people, separated by ;", None),
                   ("roles", "Roles HSV has no class for, as PersonId: Role; ...", None),
                   ("notes", "Free text, e.g. the instrument", None)],
    "Documents": [("id", "Identifier", None), ("kind", "HSV class of the document", "document kind"),
                  ("label", "Description", None), ("date", "ISO date it was created (optional)", None),
                  ("activity", "Activity that produced it", None),
                  ("derived_from", "Document it was derived from", None),
                  ("used_in", "Activity that used it", None),
                  ("authors", "Authors (HSV Author), separated by ;", None),
                  ("acquired_by", "Who acquired it (HSV Inspector)", None),
                  ("processed_by", "Who post-processed it (HSV PostProcesser)", None),
                  ("identifier", "File name, DOI or URL (optional)", None),
                  ("about", "What it documents: aggregate (default), or unit/facade identifiers", None)],
    "Units": [("id", "Identifier, e.g. Unit_417", None), ("label", "Description", None),
              ("floor_kind", "HSTO class of its floors (optional)", "floor kind"),
              ("floor_label", "Description of the floors", None),
              ("floor_accessible", "Were the floors accessible for survey?", "yes/no"),
              ("floor_note", "What is known about them and how", None)],
    "Facades": [("id", "Identifier, e.g. Facade_417a", None), ("unit", "Identifier of its unit", None),
                ("label", "Description", None), ("masonry_type", "Name of its masonry type (MasonryTypes)", None)],
    "Photos": [("id", "Identifier", None),
               ("facade", "Facades it shows, separated by ;; empty means the aggregate", None),
               ("file", "File path, relative to the photo folder", None),
               ("label", "Caption (optional)", None),
               ("taken", "ISO date and time", None), ("activity", "Activity it was taken in", None),
               ("acquired_by", "Who took it; empty means the people of the activity", None),
               ("source", "Source and licence, for photographs taken from a publication", None)],
    "Connections": [("id", "Identifier", None), ("kind", "HSTO class", "connection kind"),
                    ("part_a", "Facade on one side", None), ("part_b", "Facade on the other side", None),
                    ("label", "Description", None),
                    ("quality", "Connection quality and its evidence", None)],
    "Openings": [("id", "Identifier", None), ("facade", "Facade it is on", None),
                 ("label", "Description", None)],
    "Vulnerabilities": [("facade", "Facade whose wall has it", None),
                        ("vulnerability", "FMO vulnerability", "vulnerability"),
                        ("evidence", "What shows it, and its source", None)],
    "MasonryTypes": [("name", "Name, the same as the IFC material", None), ("description", "Description", None)]
                    + [(p, f"MQI category: {p.replace('_', ' ')}", p) for p in PARAMETERS]
                    + [(f"evidence_{p}", f"Evidence for {p.replace('_', ' ')}", None) for p in PARAMETERS]
                    + [("evidence", "General evidence (optional)", None),
                       ("flags", "Assumptions and caveats, one per line", None),
                       ("density_kg_m3", "Mass density (HMO has no rule for it)", None),
                       ("density_note", "Source of the density", None),
                       ("reference_category", "Code category it is compared with (optional)", None),
                       ("reference_note", "Note on the comparison (optional)", None),
                       ("variant_lower", "Lower variant, as parameter=value; ... (optional)", None),
                       ("variant_upper", "Upper variant, as parameter=value; ... (optional)", None)],
    "Measured": [("masonry_type", "Masonry type", None), ("quantity", "e.g. compressive_strength_MPa", None),
                 ("value", "Value", None), ("basis", "specimen or full scale", None),
                 ("source", "Where it comes from", None)],
}
PROJECT_KEYS = [
    ("id", "Short name of the case study, e.g. castelnuovo"),
    ("title", "Title shown on the page"),
    ("description", "One paragraph describing the case study"),
    ("namespace", "Namespace of the graph's IRIs, ending in #"),
    ("centre_id", "Identifier of the historic centre (optional)"),
    ("centre_label", "Name of the historic centre (optional)"),
    ("aggregate_id", "Identifier of the aggregate"),
    ("aggregate_label", "Description of the aggregate"),
    ("technique_id", "Identifier of the traditional construction technique (optional)"),
    ("technique_label", "Name of the technique (optional)"),
    ("bim_document", "Identifier of the BIM model in Documents"),
    ("software_doi", "DOI of the software release (optional)"),
]


def split(cell):
    return [x.strip() for x in str(cell or "").split(";") if x.strip()]


def text(cell):
    if cell is None:
        return ""
    if isinstance(cell, (datetime.datetime, datetime.date)):
        return cell.isoformat()
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell).strip()


def pairs(cell, sep):
    out = {}
    for item in split(cell):
        k, _, v = item.partition(sep)
        out[k.strip()] = v.strip()
    return out


def _rows(ws):
    header = [text(c.value) for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(v not in (None, "") for v in row):
            yield {h: v for h, v in zip(header, row) if h}


def read_workbook(path):
    """(case, classification) from a filled workbook."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    s = {name: list(_rows(wb[name])) if name in wb.sheetnames else [] for name in SHEETS}
    project = {text(r["key"]): text(r["value"]) for r in s["Project"]}
    case = {"project": project,
            "organisations": {text(r["id"]): text(r["name"]) for r in s["Organisations"]},
            "people": {text(r["id"]): {"name": text(r["name"]), "affiliation": text(r["organisation"])}
                       for r in s["People"]},
            "activities": {}, "documents": {}, "units": {}, "facades": {}, "photos": {},
            "connections": {}, "openings": {}, "vulnerabilities": [], "measured": []}
    for r in s["Activities"]:
        a = {"label": text(r["label"]), "people": split(r.get("people"))}
        for k in ("date", "start", "end", "notes"):
            if text(r.get(k)):
                a[k] = text(r[k])
        if text(r.get("roles")):
            a["roles"] = pairs(r["roles"], ":")
        case["activities"][text(r["id"])] = a
    for r in s["Documents"]:
        d = {"class": text(r["kind"]), "label": text(r["label"])}
        for k in ("date", "activity", "derived_from", "used_in", "identifier"):
            if text(r.get(k)):
                d[k] = text(r[k])
        for k in ("authors", "acquired_by", "processed_by", "about"):
            if split(r.get(k)):
                d[k] = split(r[k])
        case["documents"][text(r["id"])] = d
    for r in s["Units"]:
        u = {"label": text(r["label"])}
        if text(r.get("floor_kind")):
            u["floor"] = {"kind": text(r["floor_kind"]), "label": text(r.get("floor_label")),
                          "accessible": text(r.get("floor_accessible")).lower() == "yes",
                          "note": text(r.get("floor_note"))}
        case["units"][text(r["id"])] = u
    for r in s["Facades"]:
        case["facades"][text(r["id"])] = {"unit": text(r["unit"]), "label": text(r["label"]),
                                          "masonry_type": text(r["masonry_type"])}
    for r in s["Photos"]:
        case["photos"][text(r["id"])] = {k: (split(r.get(k)) if k == "acquired_by" else text(r.get(k)))
                                         for k in ("facade", "file", "label", "taken", "activity",
                                                   "acquired_by", "source")}
    for r in s["Connections"]:
        case["connections"][text(r["id"])] = {k: text(r.get(k)) for k in
                                              ("kind", "part_a", "part_b", "label", "quality")}
    for r in s["Openings"]:
        case["openings"][text(r["id"])] = {"facade": text(r["facade"]), "label": text(r["label"])}
    for r in s["Vulnerabilities"]:
        case["vulnerabilities"].append({k: text(r.get(k)) for k in ("facade", "vulnerability", "evidence")})
    for r in s["Measured"]:
        case["measured"].append({k: text(r.get(k)) for k in ("masonry_type", "quantity", "value", "basis", "source")})

    types = {}
    for r in s["MasonryTypes"]:
        t = {"description": text(r["description"]),
             "observation": {p: text(r[p]) for p in PARAMETERS},
             "evidence": {p: text(r.get(f"evidence_{p}")) for p in PARAMETERS if text(r.get(f"evidence_{p}"))},
             "flags": [x.strip() for x in text(r.get("flags")).splitlines() if x.strip()],
             "mass_density_kg_m3": float(r["density_kg_m3"]),
             "density_note": text(r.get("density_note"))}
        if text(r.get("evidence")):
            t["evidence"]["general"] = text(r["evidence"])
        if text(r.get("reference_category")):
            t["ntc_comparison"] = {"category": text(r["reference_category"]), "note": text(r.get("reference_note"))}
        variants = {k: pairs(r.get(f"variant_{k}"), "=") for k in ("lower", "upper") if text(r.get(f"variant_{k}"))}
        if variants:
            t["variants"] = variants
        types[text(r["name"])] = t
    return case, {"types": types}


def validate(case, classification):
    """Cross-references that do not resolve, as a list of messages."""
    problems = []
    c = case
    ids = {"people": set(c["people"]), "agents": set(c["people"]) | set(c["organisations"]),
           "activities": set(c["activities"]), "documents": set(c["documents"]),
           "units": set(c["units"]), "facades": set(c["facades"]), "types": set(classification["types"])}

    def need(kind, value, where):
        if value and value not in ids[kind]:
            problems.append(f"{where}: unknown {kind[:-1]} '{value}'")

    for pid, p in c["people"].items():
        if p["affiliation"] and p["affiliation"] not in c["organisations"]:
            problems.append(f"People {pid}: unknown organisation '{p['affiliation']}'")
    for aid, a in c["activities"].items():
        for pid in a["people"]:            # people or organisations
            need("agents", pid, f"Activities {aid}")
        for pid in a.get("roles", {}):
            need("people", pid, f"Activities {aid}")
    for did, d in c["documents"].items():
        for k in ("activity", "used_in"):
            need("activities", d.get(k), f"Documents {did}")
        need("documents", d.get("derived_from"), f"Documents {did}")
        for k in ("authors", "acquired_by", "processed_by"):
            for pid in d.get(k, []):
                need("people", pid, f"Documents {did}")
    for fid, f in c["facades"].items():
        need("units", f["unit"], f"Facades {fid}")
        need("types", f["masonry_type"], f"Facades {fid}")
    for phid, ph in c["photos"].items():
        for fid in split(ph["facade"]):
            need("facades", fid, f"Photos {phid}")
        need("activities", ph["activity"], f"Photos {phid}")
    for cid, cn in c["connections"].items():
        need("facades", cn["part_a"], f"Connections {cid}")
        need("facades", cn["part_b"], f"Connections {cid}")
    for oid, o in c["openings"].items():
        need("facades", o["facade"], f"Openings {oid}")
    for v in c.get("vulnerabilities", []):
        need("facades", v["facade"], "Vulnerabilities")
        if v["vulnerability"] not in CHOICES["vulnerability"]:
            problems.append(f"Vulnerabilities: '{v['vulnerability']}' is not an FMO vulnerability")
    for name, t in classification["types"].items():
        for p in PARAMETERS:
            if t["observation"][p] not in CHOICES[p]:
                problems.append(f"MasonryTypes {name}: '{t['observation'][p]}' is not a category of {p}")
    if c["project"].get("bim_document") and c["project"]["bim_document"] not in ids["documents"]:
        problems.append(f"Project: bim_document '{c['project']['bim_document']}' is not in Documents")
    return problems


def write_workbook(path, case=None, classification=None):
    """Writes the records to a workbook; with none, the empty template."""
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    lists = wb.create_sheet("_lists")
    for i, (key, values) in enumerate(CHOICES.items(), start=1):
        lists.cell(1, i, key)
        for j, v in enumerate(values, start=2):
            lists.cell(j, i, v)
    lists.sheet_state = "hidden"
    list_col = {key: i for i, key in enumerate(CHOICES, start=1)}

    head_fill = PatternFill("solid", fgColor="E9E4DA")
    rows = _records_as_rows(case, classification) if case else {}
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        for j, (col, desc, choice) in enumerate(cols, start=1):
            cell = ws.cell(1, j, col)
            cell.font, cell.fill = Font(bold=True), head_fill
            cell.comment = Comment(desc, "openBIMtoFEM")
            ws.column_dimensions[cell.column_letter].width = 16 if col in ("id", "key", "kind") else 28
            if choice:
                letter = openpyxl.utils.get_column_letter(list_col[choice])
                dv = DataValidation(type="list", allow_blank=True,
                                    formula1=f"=_lists!${letter}$2:${letter}${len(CHOICES[choice]) + 1}")
                ws.add_data_validation(dv)
                dv.add(f"{cell.column_letter}2:{cell.column_letter}500")
        ws.freeze_panes = "A2"
        if name == "Project" and not case:
            for i, (k, desc) in enumerate(PROJECT_KEYS, start=2):
                ws.cell(i, 1, k).comment = Comment(desc, "openBIMtoFEM")
        for i, row in enumerate(rows.get(name, []), start=2):
            for j, (col, _, _) in enumerate(cols, start=1):
                v = row.get(col, "")
                c = ws.cell(i, j, v if v != "" else None)
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.move_sheet("_lists", offset=len(wb.sheetnames))
    wb.save(path)


def _records_as_rows(case, classification):
    j = lambda xs: "; ".join(xs)
    rows = {"Project": [{"key": k, "value": case["project"].get(k, "")} for k, _ in PROJECT_KEYS],
            "Organisations": [{"id": k, "name": v} for k, v in case["organisations"].items()],
            "People": [{"id": k, "name": p["name"], "organisation": p["affiliation"]} for k, p in case["people"].items()],
            "Activities": [], "Documents": [], "Units": [], "Facades": [], "Photos": [],
            "Connections": [], "Openings": [], "Vulnerabilities": [], "MasonryTypes": [], "Measured": []}
    for k, a in case["activities"].items():
        rows["Activities"].append({"id": k, "label": a["label"], "date": a.get("date", ""), "start": a.get("start", ""),
                                   "end": a.get("end", ""), "people": j(a["people"]),
                                   "roles": j(f"{p}: {r}" for p, r in a.get("roles", {}).items()),
                                   "notes": a.get("notes", "")})
    for k, d in case["documents"].items():
        rows["Documents"].append({"id": k, "kind": d["class"], "label": d["label"], "date": d.get("date", ""),
                                  "activity": d.get("activity", ""), "derived_from": d.get("derived_from", ""),
                                  "used_in": d.get("used_in", ""), "authors": j(d.get("authors", [])),
                                  "acquired_by": j(d.get("acquired_by", [])),
                                  "processed_by": j(d.get("processed_by", [])),
                                  "identifier": d.get("identifier", ""), "about": j(d.get("about", []))})
    for k, u in case["units"].items():
        f = u.get("floor", {})
        rows["Units"].append({"id": k, "label": u["label"], "floor_kind": f.get("kind", ""),
                              "floor_label": f.get("label", ""),
                              "floor_accessible": ("yes" if f.get("accessible") else "no") if f else "",
                              "floor_note": f.get("note", "")})
    for k, f in case["facades"].items():
        rows["Facades"].append({"id": k, "unit": f["unit"], "label": f["label"], "masonry_type": f["masonry_type"]})
    for k, ph in case["photos"].items():
        rows["Photos"].append({**ph, "id": k, "acquired_by": j(ph.get("acquired_by", []))})
    for k, cn in case["connections"].items():
        rows["Connections"].append({"id": k, **cn})
    for k, o in case["openings"].items():
        rows["Openings"].append({"id": k, **o})
    for name, t in classification["types"].items():
        ev = t.get("evidence", {})
        if isinstance(ev, str):
            ev = {"general": ev}
        row = {"name": name, "description": t["description"], **t["observation"],
               **{f"evidence_{p}": ev.get(p, "") for p in PARAMETERS},
               "evidence": ev.get("general", ""), "flags": "\n".join(t.get("flags", [])),
               "density_kg_m3": t["mass_density_kg_m3"], "density_note": t.get("density_note", ""),
               "reference_category": (t.get("ntc_comparison") or {}).get("category", ""),
               "reference_note": (t.get("ntc_comparison") or {}).get("note", "")}
        for k, v in t.get("variants", {}).items():
            row[f"variant_{k}"] = j(f"{p}={x}" for p, x in v.items())
        rows["MasonryTypes"].append(row)
    rows["Vulnerabilities"] = list(case.get("vulnerabilities", []))
    rows["Measured"] = list(case.get("measured", []))
    return rows
